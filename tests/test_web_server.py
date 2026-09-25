"""Tests for embedded WebServer routes and warehouse API."""

from __future__ import annotations

import os
from unittest.mock import MagicMock

import pytest
from aiohttp.test_utils import TestClient, TestServer

from kolobot.warehouse_db import WarehouseDB
from kolobot.web_server import WebServer


@pytest.fixture
def warehouse_env(tmp_path):
    db_path = str(tmp_path / "warehouse.db")
    db = WarehouseDB(db_path=db_path)
    db.init_db()

    fs = MagicMock()
    fs._base = str(tmp_path)

    vs = MagicMock()
    vs.list_all_ids.return_value = []

    return db, fs, vs


@pytest.mark.asyncio
async def test_web_server_index(warehouse_env):
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/")
        assert resp.status == 200
        text = await resp.text()
        assert "kolobot" in text
        assert "edit-modal" in text
        assert "edit-new-value" in text
        assert "edit-comment" in text
        assert "openEditModal" in text
        assert "submitEditField" in text
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_items_empty(warehouse_env):
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/api/warehouse/items")
        assert resp.status == 200
        data = await resp.json()
        assert data == []
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_documents_empty(warehouse_env):
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/api/warehouse/documents")
        assert resp.status == 200
        data = await resp.json()
        assert data == []
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_documents_with_doc_type(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(
        filename="test.xlsx", file_type="excel", doc_type="НАКЛАДНА",
    )
    item_id = db.add_item(name="Болт М8", sku="12345")
    db.add_transaction(
        item_id=item_id, document_id=doc_id,
        operation_type="income", quantity=10,
        source_row="Рядок 3: 12345 | Болт М8",
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/api/warehouse/documents")
        data = await resp.json()
        assert len(data) == 1
        assert data[0]["doc_type"] == "НАКЛАДНА"
        assert data[0]["transaction_count"] == 1

        resp2 = await client.get(f"/api/warehouse/documents/{doc_id}/impact")
        impacts = await resp2.json()
        assert len(impacts) == 1
        assert impacts[0]["source_row"] == "Рядок 3: 12345 | Болт М8"

        resp3 = await client.get(f"/api/warehouse/items/{item_id}/transactions")
        txs = await resp3.json()
        assert len(txs) == 1
        assert txs[0]["doc_type"] == "НАКЛАДНА"
        assert txs[0]["source_row"] == "Рядок 3: 12345 | Болт М8"
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_delete_document(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="del.xlsx", file_type="excel")

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.delete(f"/api/warehouse/documents/{doc_id}")
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True

        resp2 = await client.delete(f"/api/warehouse/documents/{doc_id}")
        assert resp2.status == 404
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_documents_panel_renders_queued_status_column(warehouse_env):
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/")
        html = await resp.text()
        assert ">Статус<" in html
        assert "⏳ В черзі" in html
        assert "Документ у черзі на розпізнавання" in html
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_queued_document_exposes_status_and_preview(warehouse_env, tmp_path):
    db, fs, vs = warehouse_env
    image_path = tmp_path / "queued.jpg"
    image_path.write_bytes(b"\xff\xd8queued-image")
    doc_id = db.add_document(
        filename="queued.jpg",
        file_type="photo",
        file_path=str(image_path),
        status="queued",
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/api/warehouse/documents")
        data = await resp.json()
        assert len(data) == 1
        assert data[0]["status"] == "queued"
        assert data[0]["file_path"] == str(image_path)

        view = await client.get(f"/api/warehouse/documents/{doc_id}/view")
        assert view.status == 200
        assert await view.read() == b"\xff\xd8queued-image"

        ocr = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        ocr_data = await ocr.json()
        assert ocr_data["status"] == "queued"
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_completed_document_reports_completed_status(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="done.jpg", file_type="photo", raw_text="текст")

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        ocr = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        data = await ocr.json()
        assert data["status"] == "completed"
        assert data["raw_text"] == "текст"
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_web_server_processing_status_badges(warehouse_env):
    """Issue #13: Web UI exposes processing_ocr and processing_emb statuses and badges."""
    db, fs, vs = warehouse_env
    doc1 = db.add_document(filename="d1.jpg", file_type="photo", status="processing_ocr")
    doc2 = db.add_document(filename="d2.jpg", file_type="photo", status="processing_emb")

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/api/warehouse/documents")
        docs = await resp.json()
        docs_by_id = {d["id"]: d for d in docs}
        assert docs_by_id[doc1]["status"] == "processing_ocr"
        assert docs_by_id[doc2]["status"] == "processing_emb"

        html_resp = await client.get("/")
        html = await html_resp.text()
        assert "🔄 Розпізнавання" in html
        assert "🧠 Embeddings" in html
    finally:
        await client.close()
        db.close()



@pytest.mark.asyncio
async def test_api_document_ocr(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(
        filename="scan1.jpg",
        file_type="photo",
        doc_type="НАКЛАДНА",
        doc_number="0002143",
        doc_date="03.08.2026",
        raw_text="НАКЛАДНА № 0002143\nКОНДЕНСАТОР КЕРАМІЧНИЙ 0.1мФ 50В",
    )
    item_id = db.add_item(name="КОНДЕНСАТОР КЕРАМІЧНИЙ 0.1мФ 50В", sku="288410091721")
    db.add_transaction(
        item_id=item_id,
        document_id=doc_id,
        operation_type="income",
        quantity=250,
        doc_number="0002143",
        doc_date="03.08.2026",
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        assert resp.status == 200
        data = await resp.json()
        assert data["doc_number"] == "0002143"
        assert data["doc_date"] == "03.08.2026"
        assert data["doc_type"] == "НАКЛАДНА"
        assert "НАКЛАДНА № 0002143" in data["raw_text"]
        assert len(data["impact"]) == 1
        assert data["impact"][0]["name"] == "КОНДЕНСАТОР КЕРАМІЧНИЙ 0.1мФ 50В"
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_document_ocr_from_txt_file(warehouse_env, tmp_path):
    db, fs, vs = warehouse_env
    # Document has empty raw_text in db, but .txt file exists
    txt_file = tmp_path / "abc1234567.txt"
    txt_file.write_text("ВИМОГА № 0000555\nТЕКСТ З ФАЙЛУ ДИСКУ", encoding="utf-8")
    fs.read_text.return_value = "ВИМОГА № 0000555\nТЕКСТ З ФАЙЛУ ДИСКУ"

    doc_id = db.add_document(
        filename="abc1234567.jpg",
        file_type="photo",
        file_path=str(tmp_path / "abc1234567.jpg"),
        doc_type="ВИМОГА",
        doc_number="0000555",
        doc_date="15.08.2026",
        raw_text="",  # Empty in DB
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        assert resp.status == 200
        data = await resp.json()
        assert data["doc_number"] == "0000555"
        assert data["doc_date"] == "15.08.2026"
        assert data["doc_type"] == "ВИМОГА"
        assert "ВИМОГА № 0000555" in data["raw_text"]
        assert "ТЕКСТ З ФАЙЛУ ДИСКУ" in data["raw_text"]
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_edit_item_metadata_post(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="items.xlsx", file_type="excel")
    item_id = db.add_item(name="Резистор 10к", sku="RES-10K", unit="шт", supplier="ПромЕлектро", notes="0.25Вт")
    db.add_transaction(item_id=item_id, document_id=doc_id, operation_type="income", quantity=100)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # Edit name via POST /api/warehouse/items/{id}/edit
        payload = {"field": "name", "value": "Резистор 10кОм 0.5Вт", "comment": "Уточнення потужності"}
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json=payload)
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True
        assert data["item"]["name"] == "Резистор 10кОм 0.5Вт"
        assert data["old_value"] == "Резистор 10к"
        assert data["new_value"] == "Резистор 10кОм 0.5Вт"
        assert "Змінено [Найменування]: 'Резистор 10к' → 'Резистор 10кОм 0.5Вт'" in data["source_row"]

        # Verify DB transactions
        txs = db.get_item_transactions(item_id)
        assert len(txs) == 2
        assert txs[1]["quantity"] == 0.0
        assert txs[1]["doc_type"] == "РУЧНЕ_КОРИГУВАННЯ"

        # Verify items list endpoint shows updated name with same balance
        items_resp = await client.get("/api/warehouse/items")
        items = await items_resp.json()
        assert items[0]["name"] == "Резистор 10кОм 0.5Вт"
        assert items[0]["balance"] == 100.0
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_edit_item_metadata_patch(warehouse_env):
    db, fs, vs = warehouse_env
    item_id = db.add_item(name="Кабель ВВГ", sku="OLD-SKU")
    doc_id = db.add_document(filename="init.xlsx", file_type="excel")
    db.add_transaction(item_id=item_id, document_id=doc_id, operation_type="income", quantity=25)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # Edit SKU via PATCH /api/warehouse/items/{id}
        payload = {"field": "sku", "value": "CAB-VVG-3x2.5"}
        resp = await client.patch(f"/api/warehouse/items/{item_id}", json=payload)
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True
        assert data["item"]["sku"] == "CAB-VVG-3x2.5"
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_edit_item_errors(warehouse_env):
    db, fs, vs = warehouse_env
    item_id = db.add_item(name="Дріт мідний")

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 404 for nonexistent item
        resp = await client.post("/api/warehouse/items/99999/edit", json={"field": "name", "value": "Test"})
        assert resp.status == 404
        data = await resp.json()
        assert "error" in data

        # 400 for invalid JSON / empty body
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", data="not json", headers={"Content-Type": "application/json"})
        assert resp.status == 400

        # 400 for missing field
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json={"value": "Test"})
        assert resp.status == 400

        # 400 for disallowed field
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json={"field": "unsupported", "value": "Val"})
        assert resp.status == 400
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_adjust_item_quantity_success_income(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="init.xlsx", file_type="excel")
    item_id = db.add_item(name="Кабель ВВГ", sku="CAB-01", unit="м")
    db.add_transaction(item_id=item_id, document_id=doc_id, operation_type="income", quantity=50)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # Increase balance from 50 to 75 (delta = +25)
        payload = {"field": "balance", "value": 75, "comment": "Інвентаризація"}
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json=payload)
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True
        assert data["old_balance"] == 50.0
        assert data["target_quantity"] == 75.0
        assert data["delta"] == 25.0
        assert data["operation_type"] == "income"
        assert data["transaction_id"] is not None
        assert "Змінено 'Кількість': 50 → 75 (Інвентаризація)" in data["source_row"]

        # Verify item list returns updated balance
        items_resp = await client.get("/api/warehouse/items")
        items = await items_resp.json()
        assert items[0]["balance"] == 75.0
        assert items[0]["total_income"] == 75.0
        assert items[0]["total_expense"] == 0.0

        # Verify transactions endpoint returns adjustment transaction with badge info
        txs_resp = await client.get(f"/api/warehouse/items/{item_id}/transactions")
        txs = await txs_resp.json()
        assert len(txs) == 2
        assert txs[1]["doc_type"] == "РУЧНЕ_КОРИГУВАННЯ"
        assert txs[1]["filename"] == "Ручне редагування (Користувач)"
        assert txs[1]["quantity"] == 25.0
        assert txs[1]["running_balance"] == 75.0
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_adjust_item_quantity_success_expense(warehouse_env):
    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="init.xlsx", file_type="excel")
    item_id = db.add_item(name="Труба ПВХ", sku="TR-01", unit="м")
    db.add_transaction(item_id=item_id, document_id=doc_id, operation_type="income", quantity=100)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # Decrease balance from 100 to 70 (delta = -30)
        payload = {"field": "quantity", "value": "70", "comment": "Списання"}
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json=payload)
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True
        assert data["old_balance"] == 100.0
        assert data["target_quantity"] == 70.0
        assert data["delta"] == -30.0
        assert data["operation_type"] == "expense"
        assert data["transaction_id"] is not None

        # Verify balance reconciliation
        items_resp = await client.get("/api/warehouse/items")
        items = await items_resp.json()
        assert items[0]["balance"] == 70.0
        assert items[0]["total_income"] == 100.0
        assert items[0]["total_expense"] == 30.0
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_adjust_item_quantity_invalid_number(warehouse_env):
    db, fs, vs = warehouse_env
    item_id = db.add_item(name="Гайка М6")

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        payload = {"field": "balance", "value": "not-a-number"}
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json=payload)
        assert resp.status == 400
        data = await resp.json()
        assert "error" in data
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_adjust_item_quantity_and_excel_export_reconciliation(warehouse_env):
    import io
    import openpyxl

    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="init.xlsx", file_type="excel")
    item1_id = db.add_item(name="Кабель ВВГ", sku="CAB-01", unit="м")
    item2_id = db.add_item(name="Труба ПВХ", sku="TR-01", unit="м")
    db.add_transaction(item_id=item1_id, document_id=doc_id, operation_type="income", quantity=50)
    db.add_transaction(item_id=item2_id, document_id=doc_id, operation_type="income", quantity=100)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 1. Adjust item1 (+30 -> 80)
        resp1 = await client.post(f"/api/warehouse/items/{item1_id}/edit", json={"field": "balance", "value": 80, "comment": "Довезення"})
        assert resp1.status == 200

        # 2. Adjust item2 (-40 -> 60)
        resp2 = await client.post(f"/api/warehouse/items/{item2_id}/edit", json={"field": "quantity", "value": 60, "comment": "Списання"})
        assert resp2.status == 200

        # 3. Verify items API endpoint
        items_resp = await client.get("/api/warehouse/items")
        assert items_resp.status == 200
        items = await items_resp.json()
        cab = next(i for i in items if i["id"] == item1_id)
        pipe = next(i for i in items if i["id"] == item2_id)

        assert cab["total_income"] == 80.0
        assert cab["total_expense"] == 0.0
        assert cab["balance"] == 80.0

        assert pipe["total_income"] == 100.0
        assert pipe["total_expense"] == 40.0
        assert pipe["balance"] == 60.0

        # 4. Download Excel export and verify cell values
        export_resp = await client.get("/api/warehouse/export")
        assert export_resp.status == 200
        excel_bytes = await export_resp.read()

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        # Headers in row 1: sku, name, doc_date, income, expense, balance, unit, doc_number, supplier, notes
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 3  # Header + 2 items

        header = rows[0]
        income_idx = header.index("Прихід")
        expense_idx = header.index("Розхід")
        balance_idx = header.index("Залишок")
        name_idx = header.index("Найменування")

        cab_row = next(r for r in rows[1:] if r[name_idx] == "Кабель ВВГ")
        pipe_row = next(r for r in rows[1:] if r[name_idx] == "Труба ПВХ")

        assert cab_row[income_idx] == 80.0
        assert cab_row[expense_idx] == "" or cab_row[expense_idx] == 0.0  # exported as empty or 0
        assert cab_row[balance_idx] == 80.0

        assert pipe_row[income_idx] == 100.0
        assert pipe_row[expense_idx] == 40.0
        assert pipe_row[balance_idx] == 60.0
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_web_server_log_tab_html(warehouse_env):
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/")
        assert resp.status == 200
        text = await resp.text()
        assert "tab-logs" in text
        assert "panel-logs" in text
        assert "~/logs" in text
        assert "grep logs..." in text
        assert "lvl-info" in text
        assert "lvl-success" in text
        assert "lvl-warn" in text
        assert "lvl-err" in text
        assert "log-autoscroll" in text
        assert "clearLogsServer" in text
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_logs_and_clear(warehouse_env):
    from kolobot.log_service import LogBuffer

    db, fs, vs = warehouse_env
    log_buf = LogBuffer(maxlen=100)
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42, log_buffer=log_buf)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # Initial empty logs
        resp = await client.get("/api/logs")
        assert resp.status == 200
        data = await resp.json()
        assert "logs" in data
        assert "total_count" in data
        assert "last_id" in data

        # Add some log entries
        log_buf.add_entry(level="INFO", logger_name="kolobot.app", message="Bot started")
        log_buf.add_entry(level="WARN", logger_name="kolobot.key_pool", message="Key cooldown active")
        log_buf.add_entry(level="ERR", logger_name="kolobot.gateway", message="Gemini 500 error")

        # Fetch all
        resp = await client.get("/api/logs")
        assert resp.status == 200
        data = await resp.json()
        assert data["total_count"] == 3
        assert len(data["logs"]) == 3

        # Filter by since_id
        resp = await client.get("/api/logs?since_id=2")
        data = await resp.json()
        assert len(data["logs"]) == 1
        assert data["logs"][0]["level"] == "ERR"

        # Filter by level
        resp = await client.get("/api/logs?level=WARN")
        data = await resp.json()
        assert len(data["logs"]) == 1
        assert data["logs"][0]["level"] == "WARN"

        # Filter by query
        resp = await client.get("/api/logs?q=cooldown")
        data = await resp.json()
        assert len(data["logs"]) == 1
        assert "cooldown" in data["logs"][0]["message"].lower()

        # Clear via POST /api/logs/clear
        resp = await client.post("/api/logs/clear")
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True
        assert data["cleared_count"] == 3

        # Verify empty
        resp = await client.get("/api/logs")
        data = await resp.json()
        assert data["total_count"] == 0
        assert len(data["logs"]) == 0

        # Add one and clear via DELETE /api/logs
        log_buf.add_entry(level="INFO", logger_name="kolobot.app", message="Another log")
        del_resp = await client.delete("/api/logs")
        assert del_resp.status == 200
        data = await del_resp.json()
        assert data["success"] is True
        assert data["cleared_count"] == 1
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_logs_stream(warehouse_env):
    import asyncio
    from kolobot.log_service import LogBuffer

    db, fs, vs = warehouse_env
    log_buf = LogBuffer(maxlen=100)
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42, log_buffer=log_buf)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/api/logs/stream")
        assert resp.status == 200
        assert "text/event-stream" in resp.headers.get("Content-Type", "")

        # Read initial ping frame (: ping\n\n)
        line1 = await resp.content.readline()
        assert b"ping" in line1
        line2 = await resp.content.readline()
        assert line2 == b"\n"

        # Add a log entry and verify it's pushed to stream
        log_buf.add_entry(level="INFO", logger_name="kolobot.test", message="Stream event test")
        event_data = await resp.content.readline()
        assert b"data:" in event_data
        assert b"Stream event test" in event_data
    finally:
        await client.close()
        db.close()



@pytest.mark.asyncio
async def test_document_inspection_item_edit_integration(warehouse_env):
    """Integration test verifying editing item parameters from within document inspection context."""
    db, fs, vs = warehouse_env

    # 1. Ingest a document with OCR recognized items
    doc_id = db.add_document(
        filename="scan_invoice_101.jpg",
        file_type="photo",
        doc_type="НАКЛАДНА",
        doc_number="СФ-101",
        doc_date="28.08.2026",
        raw_text="РАХУНОК-ФАКТУРА СФ-101\n1. РЕЗИСТ0Р 10к (misrecognized OCR) - 100 шт",
    )
    # Misrecognized item created during OCR import
    item_id = db.add_item(name="РЕЗИСТ0Р 10к", sku="R-10K-MISREAD", unit="шт")
    db.add_transaction(
        item_id=item_id,
        document_id=doc_id,
        operation_type="income",
        quantity=100.0,
        doc_number="СФ-101",
        doc_date="28.08.2026",
        source_row="1. РЕЗИСТ0Р 10к - 100 шт",
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 2. Inspect document via OCR endpoint (as done in OCR side-by-side modal & expanded impact)
        ocr_resp = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        assert ocr_resp.status == 200
        ocr_data = await ocr_resp.json()
        assert len(ocr_data["impact"]) == 1
        imp = ocr_data["impact"][0]
        assert imp["item_id"] == item_id
        assert imp["name"] == "РЕЗИСТ0Р 10к"
        assert imp["sku"] == "R-10K-MISREAD"
        assert imp["quantity"] == 100.0

        # 3. User corrects OCR mistakes while inspecting the document (SKU and Name)
        edit_sku_resp = await client.post(
            f"/api/warehouse/items/{item_id}/edit",
            json={"field": "sku", "value": "R-10K-CORRECTED", "comment": "Виправлено помилку OCR"},
        )
        assert edit_sku_resp.status == 200
        sku_data = await edit_sku_resp.json()
        assert sku_data["success"] is True
        assert sku_data["item"]["sku"] == "R-10K-CORRECTED"

        edit_name_resp = await client.post(
            f"/api/warehouse/items/{item_id}/edit",
            json={"field": "name", "value": "Резистор 10кОм 0.25Вт", "comment": "Корекція назви за сканом"},
        )
        assert edit_name_resp.status == 200
        name_data = await edit_name_resp.json()
        assert name_data["success"] is True
        assert name_data["item"]["name"] == "Резистор 10кОм 0.25Вт"

        # 4. Re-inspect document OCR and Impact: the document items list reflects corrected values
        ocr_resp2 = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        ocr_data2 = await ocr_resp2.json()
        assert len(ocr_data2["impact"]) == 1
        updated_imp = ocr_data2["impact"][0]
        assert updated_imp["item_id"] == item_id
        assert updated_imp["name"] == "Резистор 10кОм 0.25Вт"
        assert updated_imp["sku"] == "R-10K-CORRECTED"

        impact_resp = await client.get(f"/api/warehouse/documents/{doc_id}/impact")
        assert impact_resp.status == 200
        impact_data = await impact_resp.json()
        assert impact_data[0]["name"] == "Резистор 10кОм 0.25Вт"
        assert impact_data[0]["sku"] == "R-10K-CORRECTED"

        # 5. Verify HTML template contains pencil edit action triggers in OCR list and document impact
        index_resp = await client.get("/")
        assert index_resp.status == 200
        html = await index_resp.text()
        assert "openEditModal" in html
        # Check that openEditModal is hooked up to document impact rows and OCR items list
        assert "loadDocumentOcr" in html
        assert "toggleDocImpact" in html
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_document_inspection_quantity_edit_integration(warehouse_env):
    """Edit item balance from document inspection context and verify OCR/impact reflect the change."""
    db, fs, vs = warehouse_env

    doc_id = db.add_document(
        filename="scan_nakladna_200.jpg",
        file_type="photo",
        doc_type="НАКЛАДНА",
        doc_number="НКЛ-200",
        doc_date="29.08.2026",
        raw_text="НАКЛАДНА НКЛ-200\nТРУБА ПВХ 32мм - 50 м",
    )
    item_id = db.add_item(name="Труба ПВХ 32мм", sku="PIPE-32", unit="м")
    db.add_transaction(
        item_id=item_id,
        document_id=doc_id,
        operation_type="income",
        quantity=50.0,
        doc_number="НКЛ-200",
        doc_date="29.08.2026",
        source_row="1. ТРУБА ПВХ 32мм - 50 м",
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 1. Inspect document — see item with balance 50
        ocr_resp = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        assert ocr_resp.status == 200
        ocr_data = await ocr_resp.json()
        assert ocr_data["impact"][0]["quantity"] == 50.0

        # 2. User corrects quantity from the document inspection view (OCR was 50, actual is 45)
        edit_resp = await client.post(
            f"/api/warehouse/items/{item_id}/edit",
            json={"field": "balance", "value": 45, "comment": "Перерахунок за сканом"},
        )
        assert edit_resp.status == 200
        data = await edit_resp.json()
        assert data["success"] is True
        assert data["old_balance"] == 50.0
        assert data["target_quantity"] == 45.0
        assert data["delta"] == -5.0
        assert data["operation_type"] == "expense"

        # 3. Verify warehouse items endpoint reflects the new balance
        items_resp = await client.get("/api/warehouse/items")
        items = await items_resp.json()
        pipe = next(i for i in items if i["id"] == item_id)
        assert pipe["balance"] == 45.0

        # 4. Verify audit trail in transactions
        txs_resp = await client.get(f"/api/warehouse/items/{item_id}/transactions")
        txs = await txs_resp.json()
        assert len(txs) == 2
        adj_tx = txs[1]
        assert adj_tx["doc_type"] == "РУЧНЕ_КОРИГУВАННЯ"
        assert adj_tx["quantity"] == 5.0
        assert adj_tx["operation_type"] == "expense"
        assert adj_tx["running_balance"] == 45.0
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_document_inspection_multi_item_edit_integration(warehouse_env):
    """Edit multiple items from a single document's inspection view."""
    db, fs, vs = warehouse_env

    doc_id = db.add_document(
        filename="scan_nakladna_300.jpg",
        file_type="photo",
        doc_type="НАКЛАДНА",
        doc_number="НКЛ-300",
        doc_date="30.08.2026",
        raw_text="НАКЛАДНА НКЛ-300\n1. Б0ЛТ М8 - 200 шт\n2. ГАЙКА М8 - 150 шт",
    )
    item1_id = db.add_item(name="Б0ЛТ М8", sku="BOLT-M8-BAD", unit="шт")
    item2_id = db.add_item(name="ГАЙКА М8", sku="NUT-M8", unit="шт")
    db.add_transaction(item_id=item1_id, document_id=doc_id, operation_type="income", quantity=200.0)
    db.add_transaction(item_id=item2_id, document_id=doc_id, operation_type="income", quantity=150.0)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 1. Inspect document — both items visible
        ocr_resp = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        ocr = await ocr_resp.json()
        assert len(ocr["impact"]) == 2

        # 2. Fix OCR mistake in item1 name (Б0ЛТ -> Болт) and SKU
        resp1 = await client.post(
            f"/api/warehouse/items/{item1_id}/edit",
            json={"field": "name", "value": "Болт М8", "comment": "OCR помилка: 0 замість о"},
        )
        assert resp1.status == 200

        resp2 = await client.post(
            f"/api/warehouse/items/{item1_id}/edit",
            json={"field": "sku", "value": "BOLT-M8-FIXED"},
        )
        assert resp2.status == 200

        # 3. Edit item2 unit
        resp3 = await client.post(
            f"/api/warehouse/items/{item2_id}/edit",
            json={"field": "unit", "value": "компл", "comment": "Комплект, не штуки"},
        )
        assert resp3.status == 200

        # 4. Re-inspect document — all edits reflected
        ocr2_resp = await client.get(f"/api/warehouse/documents/{doc_id}/ocr")
        ocr2 = await ocr2_resp.json()
        assert len(ocr2["impact"]) == 2

        bolt = next(i for i in ocr2["impact"] if i["item_id"] == item1_id)
        nut = next(i for i in ocr2["impact"] if i["item_id"] == item2_id)
        assert bolt["name"] == "Болт М8"
        assert bolt["sku"] == "BOLT-M8-FIXED"
        assert nut["unit"] == "компл"

        # 5. Verify impact endpoint also matches
        impact_resp = await client.get(f"/api/warehouse/documents/{doc_id}/impact")
        impact = await impact_resp.json()
        bolt_imp = next(i for i in impact if i["item_id"] == item1_id)
        nut_imp = next(i for i in impact if i["item_id"] == item2_id)
        assert bolt_imp["name"] == "Болт М8"
        assert bolt_imp["sku"] == "BOLT-M8-FIXED"
        assert nut_imp["unit"] == "компл"

        # 6. Verify audit transactions: 2 edits for item1 (name, sku), 1 for item2 (unit)
        txs1 = db.get_item_transactions(item1_id)
        audit_txs1 = [t for t in txs1 if t["doc_type"] == "РУЧНЕ_КОРИГУВАННЯ"]
        assert len(audit_txs1) == 2
        assert all(t["quantity"] == 0.0 for t in audit_txs1)

        txs2 = db.get_item_transactions(item2_id)
        audit_txs2 = [t for t in txs2 if t["doc_type"] == "РУЧНЕ_КОРИГУВАННЯ"]
        assert len(audit_txs2) == 1
        assert "Змінено [Од. виміру]: 'шт' → 'компл'" in audit_txs2[0]["source_row"]
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_web_server_index_min_balance_column_and_filter(warehouse_env):
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.get("/")
        assert resp.status == 200
        text = await resp.text()
        assert "Мін. залишок" in text
        assert "filter-below-min" in text
        assert "filter-below-min-count" in text
        assert "Менше мінімального залишку" in text
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_edit_min_balance_post_and_patch(warehouse_env):
    db, fs, vs = warehouse_env
    item_id = db.add_item(name="Кабель UTP", sku="UTP-5E", min_balance=0.0)
    doc_id = db.add_document(filename="init.xlsx", file_type="excel")
    db.add_transaction(item_id=item_id, document_id=doc_id, operation_type="income", quantity=50)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 1. Edit min_balance via POST /api/warehouse/items/{id}/edit
        payload = {"field": "min_balance", "value": 20, "comment": "Встановлено мінімальний залишок"}
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json=payload)
        assert resp.status == 200
        data = await resp.json()
        assert data["success"] is True
        assert data["item"]["min_balance"] == 20.0
        assert "Змінено [Мінімальний залишок]: '0' → '20'" in data["source_row"]

        # 2. Check items API returns min_balance
        items_resp = await client.get("/api/warehouse/items")
        items = await items_resp.json()
        assert len(items) == 1
        assert items[0]["min_balance"] == 20.0
        assert items[0]["balance"] == 50.0

        # 3. Edit min_balance via PATCH
        patch_payload = {"field": "min_balance", "value": "35.5"}
        resp_patch = await client.patch(f"/api/warehouse/items/{item_id}", json=patch_payload)
        assert resp_patch.status == 200
        data_patch = await resp_patch.json()
        assert data_patch["success"] is True
        assert data_patch["item"]["min_balance"] == 35.5
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_edit_min_balance_invalid_number(warehouse_env):
    db, fs, vs = warehouse_env
    item_id = db.add_item(name="Тестовий товар")

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        resp = await client.post(f"/api/warehouse/items/{item_id}/edit", json={"field": "min_balance", "value": "abc"})
        assert resp.status == 400
        data = await resp.json()
        assert "error" in data
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_min_balance_excel_export_and_import(warehouse_env, tmp_path):
    import io
    import openpyxl
    from kolobot.excel_service import parse_excel

    db, fs, vs = warehouse_env
    doc_id = db.add_document(filename="init.xlsx", file_type="excel")
    item1_id = db.add_item(name="Світильник 36W", sku="LIGHT-36", min_balance=15.0)
    db.add_transaction(item_id=item1_id, document_id=doc_id, operation_type="income", quantity=10.0)

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        # 1. Export Excel and check min_balance column
        export_resp = await client.get("/api/warehouse/export")
        assert export_resp.status == 200
        excel_bytes = await export_resp.read()

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 2
        header = rows[0]
        assert "Мін. залишок" in header
        min_idx = header.index("Мін. залишок")
        assert rows[1][min_idx] == 15.0

        # 2. Test parse_excel with minimum balance header
        test_file = tmp_path / "test_import.xlsx"
        wb_new = openpyxl.Workbook()
        ws_new = wb_new.active
        ws_new.append(["Найменування", "Номенклатурний номер", "Прихід", "Мінімальний залишок"])
        ws_new.append(["Автомат 16A", "AUTO-16", 20, 5])
        wb_new.save(str(test_file))
        wb_new.close()

        doc_type, parsed_rows = parse_excel(str(test_file))
        assert len(parsed_rows) == 1
        assert parsed_rows[0]["name"] == "Автомат 16A"
        assert parsed_rows[0]["min_balance"] == 5.0
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_api_documents_and_transactions_expose_requested_by(warehouse_env):
    """У "Документах" і в розгорнутому рядку "Складу" видно, хто затребував документ і через кого."""
    db, fs, vs = warehouse_env
    doc_id = db.add_document(
        filename="vymoha.jpg",
        file_type="photo",
        doc_type="ВИМОГА",
        raw_text="ВИМОГА № 0000215\nЧЕРЕЗ КОГО 7939 - (ПІБ)\nЗАТРЕБУВАВ: начальник служби (ПІБ)",
        requested_by="начальник служби (ПІБ)",
        requested_via="7939 - (ПІБ)",
    )
    item_id = db.add_item(name="Болт М8", sku="12345")
    db.add_transaction(
        item_id=item_id, document_id=doc_id,
        operation_type="expense", quantity=5,
    )

    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        docs = await (await client.get("/api/warehouse/documents")).json()
        assert docs[0]["requested_by"] == "начальник служби (ПІБ)"
        assert docs[0]["requested_via"] == "7939 - (ПІБ)"

        txs = await (await client.get(f"/api/warehouse/items/{item_id}/transactions")).json()
        assert txs[0]["requested_by"] == "начальник служби (ПІБ)"
        assert txs[0]["requested_via"] == "7939 - (ПІБ)"

        ocr = await (await client.get(f"/api/warehouse/documents/{doc_id}/ocr")).json()
        assert ocr["requested_by"] == "начальник служби (ПІБ)"
        assert ocr["requested_via"] == "7939 - (ПІБ)"
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_documents_table_has_requested_by_column(warehouse_env):
    """"Документи" мають однойменну колонку, "Склад" — колонку в таблиці транзакцій."""
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        html = await (await client.get("/")).text()
        docs_head = html.split('id="docs-tbody"')[0]
        assert docs_head.count("Затребував") == 1
        assert docs_head.count("Через кого") == 1
        assert "tx.requested_by" in html
        assert "tx.requested_via" in html
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_table_fields_wrap_instead_of_truncating(warehouse_env):
    """Табличні поля переносяться рядками, а не обрізаються трикрапкою."""
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        html = await (await client.get("/")).text()
        # "data.truncated" — прапорець обрізання OCR-тексту в API, не CSS-клас
        markup = html.replace("data.truncated", "")

        assert "truncate" not in markup
        for legacy_width in ("max-w-[105px]", "max-w-[160px]", "max-w-[170px]",
                             "max-w-[180px]", "max-w-[200px]", "max-w-[250px]"):
            assert legacy_width not in markup

        # перенос рядків у клітинках дозволено
        assert markup.count("break-words") >= 10

        # title-підказки, що лише компенсували обрізання, прибрані
        assert 'title="${esc(requestedBy)}"' not in html
        assert 'title="${esc(requestedVia)}"' not in html
        assert 'title="${esc(it.notes)}"' not in html
        assert 'title="${esc(imp.name)}"' not in html
        assert 'title="${esc(srcRow)}"' not in html
        assert 'title="${esc(tx.filename)}"' not in html

        # підказка з текстом помилки на бейджі статусу лишається
        assert 'title="${esc(doc.error_message)}"' in html
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_tables_use_compact_cells(warehouse_env):
    """Обидві таблиці компактизовано, а правила мають !important."""
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        html = await (await client.get("/")).text()

        # клас стоїть на обох таблицях («Склад» і «Документи»)
        assert html.count("border-collapse compact-table") == 2

        # tracking-wider прибрано саме із заголовків таблиць (у логах він лишається)
        assert html.count("uppercase tracking-wider border-b") == 0
        assert html.count("text-slate-400 text-xs font-semibold uppercase border-b border-slate-800") == 2

        # !important обов'язковий: Tailwind з CDN вставляє утиліти в <head> пізніше
        # за вбудований <style>, тому без нього правила мовчки програють px-3 / py-4
        assert ".compact-table > thead > tr > th" in html
        assert ".compact-table > tbody > tr > td:not([colspan])" in html
        assert "padding: 8px !important;" in html
        assert "font-size: 10px !important;" in html
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_documents_table_has_card_mode_markup(warehouse_env):
    """Картковий режим нижче 1000px: маркер-клас, data-label і адаптивний CSS."""
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        html = await (await client.get("/")).text()

        # маркер стоїть на обох таблицях — «Склад» доєднався у слайсі #22
        assert html.count("compact-table card-table") == 2

        # кожна клітинка рядка документа підписана текстом свого <th>
        for label in ("Превʼю", "Файл", "Тип", "Тип документу", "Статус",
                      "Дата завантаження", "№ документа", "Затребував",
                      "Через кого", "Позицій", "Дії"):
            assert f'data-label="{label}"' in html, label
        # рахунок обмежено тілом renderDocs, щоб підписи інших таблиць і CSS-селектори
        # з data-label="…" не впливали на число
        render_docs = html.split("function renderDocs(docs)")[1].split("tbody.innerHTML = html;")[0]
        assert render_docs.count(' data-label="') == 11

        # шеврон і повноширинні клітинки розгорнутих блоків підпису не отримують
        assert '<td class="py-4 px-3"><i id="doc-chevron-' in html
        assert '<td colspan="12" class="p-0">' in html  # рядок деталей — без data-label

        # перемикання таблиця/картки робить виключно CSS, без JS-розгалужень за шириною.
        # 999.98px, а не 1000px: межа мусить працювати в обидва боки —
        # при 1000px ще таблиця, при 999px уже картки.
        assert "@media (max-width: 999.98px)" in html
        assert "attr(data-label)" in html
        assert ".card-table thead { display: none; }" in html
        assert "innerWidth" not in html

        # шеврон у картці переїжджає у правий верхній кут, а не лишається порожнім рядком
        assert ".card-table > tbody > tr:not(.hidden) > td:first-child:not([colspan])" in html
        assert "white-space: nowrap;" in html  # числа й дати не рвуться посеред значення
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_warehouse_table_has_card_mode_markup(warehouse_env):
    """Слайс #22: таблиця «Склад» стає картками, олівці видимі й стоять біля значення."""
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        html = await (await client.get("/")).text()

        # кожна з дев'яти підписаних колонок складу підписується текстом свого <th>
        for label in ("Ном. номер", "Найменування", "Прихід", "Розхід", "Залишок",
                      "Мін. залишок", "Од.виміру", "Постачальник", "Примітки"):
            assert f'data-label="{label}"' in html, label

        # десята колонка — шеврон: підпису не має, як і рядок розгорнутих транзакцій
        assert '<td class="py-4 px-3"><i id="chevron-' in html
        assert '<td colspan="10" class="p-0">' in html

        # рахунок обмежено тілом renderItems: CSS-селектори з data-label="…" і підписи
        # вкладених таблиць не мають впливати на число
        render_items = html.split("function renderItems(items)")[1].split("tbody.innerHTML = html;")[0]
        assert render_items.count(' data-label="') == 9

        # олівці редагування: у картці вони opacity-1 без наведення, а justify-content
        # flex-start перебиває утиліту justify-between, тож олівець стоїть біля значення
        assert ".card-table > tbody > tr > td button { opacity: 1 !important; }" in html
        assert ".card-table > tbody > tr > td > span.flex { justify-content: flex-start !important; }" in html

        # підсвітка «нижче мінімуму»: правило картки специфічніше за утиліти Tailwind з CDN,
        # тому рамку й тло повертає окреме правило за маркером row-low
        assert "rowClass = 'row-low bg-rose-950/40" in html
        assert ".card-table > tbody > tr.row-low" in html
        assert "border-left: 4px solid #f43f5e !important;" in html

        # числові колонки складу не рвуться посеред значення
        for label in ("Прихід", "Розхід", "Залишок", "Мін. залишок"):
            assert f'.card-table > tbody > tr > td[data-label="{label}"]' in html, label
    finally:
        await client.close()
        db.close()


@pytest.mark.asyncio
async def test_nested_tables_have_card_mode_markup(warehouse_env):
    """Слайс #23: історія транзакцій і позиції документа теж стають картками."""
    db, fs, vs = warehouse_env
    server = WebServer(warehouse_db=db, file_store=fs, vector_store=vs, owner_user_id=42)
    client = TestClient(TestServer(server._app))
    await client.start_server()

    try:
        html = await (await client.get("/")).text()

        # історія транзакцій: 10 підписів усередині самої таблиці
        tx_table = html.split('<table class="w-full card-table">')[1].split("</table>")[0]
        for label in ("Дата", "Тип", "Тип док.", "Кількість", "№ накл.",
                      "Залишок", "Документ", "Затребував", "Через кого", "Джерело"):
            assert f'data-label="{label}"' in tx_table, label
        assert tx_table.count(' data-label="') == 10

        # позиції документа: два окремі рендери — для фото й для Excel —
        # з однаковим набором із 6 колонок в обох
        item_tables = html.split('<table class="w-full text-xs card-table">')[1:]
        assert len(item_tables) == 2
        for body in (chunk.split("</table>")[0] for chunk in item_tables):
            for label in ("Ном. номер", "Найменування", "Тип", "Кількість", "Од.", "Джерело"):
                assert f'data-label="{label}"' in body, label
            assert body.count(' data-label="') == 6

        # перша клітинка вкладеної таблиці — змістовна колонка, а не шеврон,
        # тож правило правого верхнього кута з #21 для неї скасовується
        assert ".card-table .card-table > tbody > tr:not(.hidden) > td:first-child:not([colspan])" in html

        # порожній стан не став таблицею — це звичайний абзац, який картковий режим не чіпає
        assert '\'<p class="text-slate-500 py-2">Немає транзакцій.</p>\'' in html

        # дати, кількості й номери накладних не рвуться посеред значення
        for label in ("Дата", "Кількість", "№ накл."):
            assert f'.card-table > tbody > tr > td[data-label="{label}"]' in html, label
    finally:
        await client.close()
        db.close()



