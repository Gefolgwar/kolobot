"""ExcelService: parse and export warehouse Excel files."""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side


HEADER_MAP = {
    "найменування": "name",
    "найменування та технічні характеристики": "name",
    "найменування, сорт, розмір матеріалу": "name",
    "дата": "doc_date",
    "прихід": "income",
    "розхід": "expense",
    "залишок": "balance",
    "од.виміру": "unit",
    "од. виміру": "unit",
    "одиниці виміру": "unit",
    "од вим": "unit",
    "№ накл.": "doc_number",
    "№ накл": "doc_number",
    "номер накладної": "doc_number",
    "постачальник": "supplier",
    "примітки": "notes",
    "номенклатурний номер": "sku",
    "тип документу": "doc_type",
    "відпущено": "expense",
    "затребувано": "income",
    "кількість": "income",
    "ціна": "price",
    "сума": "total",
}

DOC_TYPE_KEYWORDS = {
    "вимога": "ВИМОГА",
    "накладна": "НАКЛАДНА",
}


def _normalize_header(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip().lower()


def detect_doc_type(ws: Any, max_rows: int = 10) -> str:
    for row in ws.iter_rows(max_row=max_rows, values_only=True):
        for cell_val in row:
            if cell_val is None:
                continue
            text = str(cell_val).strip().lower()
            for keyword, label in DOC_TYPE_KEYWORDS.items():
                if keyword in text:
                    return label
    return ""


def parse_excel(file_path: str) -> tuple[str, List[Dict[str, Any]]]:
    wb = openpyxl.load_workbook(file_path)
    ws = wb.worksheets[0]

    doc_type = detect_doc_type(ws)

    header_row_idx = None
    col_map: Dict[int, str] = {}

    for row_idx, row in enumerate(ws.iter_rows(max_row=10, values_only=False), start=1):
        for cell in row:
            normalized = _normalize_header(cell.value)
            if normalized in HEADER_MAP:
                header_row_idx = row_idx
                break
        if header_row_idx:
            break

    if header_row_idx is None:
        wb.close()
        return doc_type, []

    for cell in ws[header_row_idx]:
        normalized = _normalize_header(cell.value)
        if normalized in HEADER_MAP:
            col_map[cell.column] = HEADER_MAP[normalized]

    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row_idx + 1, values_only=False), start=header_row_idx + 1):
        record: Dict[str, Any] = {
            "name": "", "sku": "", "doc_date": "", "income": None,
            "expense": None, "balance": None, "unit": "", "doc_number": "",
            "supplier": "", "notes": "", "doc_type": "", "price": None, "total": None,
        }
        has_data = False
        for cell in row:
            field = col_map.get(cell.column)
            if field and cell.value is not None:
                record[field] = cell.value
                if field == "name" and str(cell.value).strip():
                    has_data = True
                if field in ("income", "expense", "balance") and cell.value:
                    has_data = True

        if has_data and record["name"]:
            record["name"] = str(record["name"]).strip()
            record["sku"] = str(record.get("sku", "") or "").strip()
            record["unit"] = str(record.get("unit", "") or "").strip()
            record["supplier"] = str(record.get("supplier", "") or "").strip()
            record["notes"] = str(record.get("notes", "") or "").strip()
            record["doc_number"] = str(record.get("doc_number", "") or "").strip()
            record["doc_date"] = str(record.get("doc_date", "") or "").strip()
            record["doc_type"] = str(record.get("doc_type", "") or "").strip()
            record["row_index"] = row_idx

            source_parts = []
            if record["sku"]:
                source_parts.append(record["sku"])
            source_parts.append(record["name"])
            if record["unit"]:
                source_parts.append(record["unit"])
            record["source_row"] = f"Рядок {row_idx}: " + " | ".join(source_parts)

            for num_field in ("income", "expense", "balance", "price", "total"):
                val = record[num_field]
                if val is not None:
                    try:
                        record[num_field] = float(val)
                    except (ValueError, TypeError):
                        record[num_field] = None

            if not record.get("doc_type") and doc_type:
                record["doc_type"] = doc_type

            items.append(record)

    wb.close()
    return doc_type, items


def export_excel(items: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Склад"

    headers = [
        "Номенклатурний номер", "Найменування", "Дата", "Прихід",
        "Розхід", "Залишок", "Од.виміру", "№ накл.",
        "Постачальник", "Примітки",
    ]
    field_keys = [
        "sku", "name", "last_doc_date", "total_income",
        "total_expense", "balance", "unit", "last_doc_number",
        "supplier", "notes",
    ]

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_font = Font(bold=True, size=11)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    for row_idx, item in enumerate(items, start=2):
        for col_idx, key in enumerate(field_keys, start=1):
            val = item.get(key, "")
            if val is None:
                val = ""
            if isinstance(val, float) and val == 0.0:
                val = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    col_widths = [22, 45, 14, 10, 10, 10, 12, 12, 25, 20]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def preview_excel(file_path: str, max_rows: int = 100) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(file_path, read_only=True)
    ws = wb.worksheets[0]
    headers: List[str] = []
    rows: List[List[str]] = []
    row_numbers: List[int] = []
    truncated = False
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [str(c) if c is not None else "" for c in row]
        if row_idx == 0 or (not headers and any(vals)):
            if not headers:
                headers = vals
                continue
        if any(v.strip() for v in vals):
            rows.append(vals)
            row_numbers.append(row_idx + 1)
        if len(rows) >= max_rows:
            truncated = True
            break
    wb.close()
    return {"headers": headers, "rows": rows, "row_numbers": row_numbers, "total_rows": len(rows), "truncated": truncated}
